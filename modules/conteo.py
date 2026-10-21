"""Conteo inicial del inventario con una planilla Excel (Paso 7 del módulo de alertas).

Problema: `stock_actual` hoy es la suma de todo lo comprado desde enero, no lo
que hay en el estante. Antes de usar las alertas hay que contar los repuestos
y corregir el stock.

Flujo:
    1. PlanillaConteo.generar(...)        -> Excel con todos los repuestos activos
                                             y la columna "Cantidad contada" vacía.
    2. Se recorre el taller y se completa la planilla (se puede hacer de a partes:
       las filas vacías se saltean).
    3. PlanillaConteo.leer(...)            -> filas leídas + errores de formato.
    4. ServicioConteoInicial.aplicar(...)  -> valida TODO; si hay algún error no
       guarda nada; si no, registra los AJUSTES, recalcula los umbrales y revisa
       las alertas.

Responsabilidades separadas (SRP, Unidad 3):
    - PlanillaConteo: solo sabe del formato del archivo (columnas, tipos).
    - ServicioConteoInicial: coordina el conteo con el gestor, los umbrales y
      las alertas. No sabe nada de Excel.
"""
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path
from typing import Callable, Dict, List, Optional, Union

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from modules.alertas import MotorAlertas, ServicioParametrosStock
from modules.modelos import Alerta, Repuesto

COLUMNAS = (
    "ID",
    "Código",
    "Marca",
    "Categoría",
    "Especificaciones",
    "Unidad",
    "Stock en el sistema",
    "Cantidad contada",
)
NOMBRE_HOJA = "Conteo"
_COL_ID, _COL_CODIGO, _COL_CONTADO = 0, 1, 7


# ---------------------------------------------------------------- Excepciones y datos

class PlanillaConteoInvalidaError(ValueError):
    """El archivo no es una planilla de conteo (falta la hoja o cambiaron las columnas)."""


@dataclass(frozen=True)
class FilaConteo:
    """Un renglón de la planilla ya leído. `fila` es el número de fila en Excel."""
    fila: int
    repuesto_id: int
    codigo: str
    cantidad_contada: Optional[float]  # None = todavía no se contó


@dataclass(frozen=True)
class ErrorConteo:
    """Un problema en un renglón, con lo necesario para encontrarlo en la planilla."""
    fila: int
    codigo: str
    mensaje: str


@dataclass
class LecturaPlanilla:
    filas: List[FilaConteo] = field(default_factory=list)
    errores: List[ErrorConteo] = field(default_factory=list)


@dataclass
class ResultadoConteo:
    """Qué pasó al aplicar la planilla."""
    ajustados: int = 0           # el stock cambió
    sin_diferencia: int = 0      # lo contado coincidía con el sistema
    sin_contar: int = 0          # filas vacías: se saltearon
    umbrales_recalculados: int = 0
    alertas_nuevas: List[Alerta] = field(default_factory=list)
    errores: List[ErrorConteo] = field(default_factory=list)

    @property
    def aplicado(self) -> bool:
        """True si se guardó el conteo (no hubo ningún error)."""
        return not self.errores


# ---------------------------------------------------------------- Planilla (formato Excel)

class PlanillaConteo:
    """Escribe y lee la planilla de conteo. Solo conoce el formato del archivo."""

    def generar(self, ruta: Union[str, Path], repuestos: List[Repuesto],
                marcas: Dict[int, str], categorias: Dict[int, str]) -> int:
        """Crea el Excel con un renglón por repuesto. Devuelve cuántos renglones escribió.

        Orden: categoría, marca y código, para recorrer el taller estante por estante.
        Pre: `marcas` y `categorias` traducen id -> nombre.
        """
        libro = Workbook()
        hoja = libro.active
        hoja.title = NOMBRE_HOJA
        hoja.append(COLUMNAS)

        ordenados = sorted(
            repuestos,
            key=lambda r: (categorias.get(r.categoria_id, ""), marcas.get(r.marca_id, ""), r.codigo),
        )
        for rep in ordenados:
            hoja.append((
                rep.id,
                rep.codigo,
                marcas.get(rep.marca_id, ""),
                categorias.get(rep.categoria_id, ""),
                rep.especificaciones or "",
                rep.unidad_medida,
                rep.stock_actual,
                None,
            ))

        self._dar_formato(hoja, len(ordenados))
        libro.save(ruta)
        return len(ordenados)

    @staticmethod
    def _dar_formato(hoja, cantidad: int) -> None:
        for celda in hoja[1]:
            celda.font = Font(bold=True)
            celda.alignment = Alignment(wrap_text=True, vertical="center")
        for letra, ancho in zip("ABCDEFGH", (6, 28, 14, 18, 30, 9, 12, 12)):
            hoja.column_dimensions[letra].width = ancho
        hoja.freeze_panes = "A2"
        if cantidad == 0:
            return
        # La única columna para completar va resaltada y solo acepta números >= 0
        # (prevención de errores: se avisa en el momento, en el mismo Excel).
        rango = f"H2:H{cantidad + 1}"
        amarillo = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
        for (celda,) in hoja[rango]:
            celda.fill = amarillo
        validacion = DataValidation(
            type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
            showErrorMessage=True, errorTitle="Cantidad inválida",
            error="Escribí la cantidad que hay en el estante (un número, 0 o más).",
        )
        validacion.add(rango)
        hoja.add_data_validation(validacion)

    def leer(self, ruta: Union[str, Path]) -> LecturaPlanilla:
        """Lee la planilla completada.

        Pre: tiene la hoja "Conteo" con las columnas de COLUMNAS, si no PlanillaConteoInvalidaError.
        Post: una FilaConteo por renglón con ID; un ErrorConteo por cada cantidad
              que no es un número. Acepta coma decimal escrita como texto ("2,5").
        """
        libro = load_workbook(ruta, data_only=True)
        if NOMBRE_HOJA not in libro.sheetnames:
            raise PlanillaConteoInvalidaError(f'El archivo no tiene la hoja "{NOMBRE_HOJA}".')
        hoja = libro[NOMBRE_HOJA]
        encabezado = tuple(c.value for c in hoja[1][:len(COLUMNAS)])
        if encabezado != COLUMNAS:
            raise PlanillaConteoInvalidaError(
                "Las columnas no son las de la planilla de conteo: generá una nueva y no cambies los títulos."
            )

        lectura = LecturaPlanilla()
        for numero, valores in enumerate(hoja.iter_rows(min_row=2, values_only=True), start=2):
            if all(v is None for v in valores):
                continue
            codigo = str(valores[_COL_CODIGO] or "")
            try:
                repuesto_id = int(valores[_COL_ID])
            except (TypeError, ValueError):
                lectura.errores.append(ErrorConteo(numero, codigo, "Falta el ID del repuesto o no es un número."))
                continue
            try:
                contado = self._a_numero(valores[_COL_CONTADO])
            except ValueError:
                lectura.errores.append(ErrorConteo(
                    numero, codigo, f'La cantidad contada "{valores[_COL_CONTADO]}" no es un número.'
                ))
                continue
            lectura.filas.append(FilaConteo(numero, repuesto_id, codigo, contado))
        return lectura

    @staticmethod
    def _a_numero(valor) -> Optional[float]:
        if valor is None or (isinstance(valor, str) and not valor.strip()):
            return None
        if isinstance(valor, bool):
            raise ValueError(valor)
        if isinstance(valor, (int, float)):
            return float(valor)
        return float(str(valor).strip().replace(",", "."))


# ---------------------------------------------------------------- Servicio

class ServicioConteoInicial:
    """Aplica una planilla de conteo: todo o nada.

    Colaboradores (se reciben, no se crean: inversión de dependencias, Unidad 3):
      - gestor: GestorInventarioBD (repuestos y movimientos).
      - motor: MotorAlertas (revisa las alertas al final).
      - fabrica_parametros: arma el ServicioParametrosStock DESPUÉS de guardar
        los ajustes, para que el período llegue hasta la fecha del conteo.
    """

    def __init__(self, gestor, motor: MotorAlertas,
                 fabrica_parametros: Callable[..., ServicioParametrosStock] = ServicioParametrosStock.desde_historial):
        self._gestor = gestor
        self._motor = motor
        self._fabrica_parametros = fabrica_parametros

    def validar(self, filas: List[FilaConteo], fecha: str) -> List[ErrorConteo]:
        """Revisa todas las filas SIN guardar nada.

        Detecta: ID repetido, repuesto inexistente, código que no corresponde al
        ID (la planilla se desordenó o se editó) y cantidades inválidas.
        """
        errores = []
        vistos = {}
        for f in filas:
            if f.repuesto_id in vistos:
                errores.append(ErrorConteo(f.fila, f.codigo, f"El ID {f.repuesto_id} está repetido (también en la fila {vistos[f.repuesto_id]})."))
                continue
            vistos[f.repuesto_id] = f.fila
            rep = self._gestor.obtener_repuesto_id(f.repuesto_id)
            if rep is None:
                errores.append(ErrorConteo(f.fila, f.codigo, f"No existe un repuesto con ID {f.repuesto_id}."))
                continue
            if rep.codigo != f.codigo:
                errores.append(ErrorConteo(f.fila, f.codigo, f'El ID {f.repuesto_id} corresponde al código "{rep.codigo}", no a "{f.codigo}".'))
                continue
            if f.cantidad_contada is not None:
                try:
                    rep.ajuste_por_conteo(f.cantidad_contada, fecha)  # solo valida: no guarda
                except ValueError as error:
                    errores.append(ErrorConteo(f.fila, f.codigo, str(error)))
        return errores

    def aplicar(self, lectura: LecturaPlanilla, fecha: Optional[str] = None) -> ResultadoConteo:
        """Pre: ninguna (los problemas se devuelven en `errores`, no se lanzan).
        Post: si hay algún error no se guarda nada (`aplicado` = False);
              si no, cada fila con cantidad deja stock_actual = lo contado,
              se recalculan los umbrales de todos y se revisan las alertas
              solo de los repuestos contados.
        """
        fecha = fecha or date.today().isoformat()
        errores = lectura.errores + self.validar(lectura.filas, fecha)
        if errores:
            return ResultadoConteo(errores=sorted(errores, key=lambda e: e.fila))

        resultado = ResultadoConteo()
        contados = []
        for f in lectura.filas:
            if f.cantidad_contada is None:
                resultado.sin_contar += 1
                continue
            contados.append(f.repuesto_id)
            if self._gestor.registrar_conteo(f.repuesto_id, f.cantidad_contada, fecha).cantidad == 0:
                resultado.sin_diferencia += 1
            else:
                resultado.ajustados += 1

        if contados:
            resultado.umbrales_recalculados = self._fabrica_parametros(self._gestor).recalcular_todos()
            # Solo se revisan los repuestos CONTADOS: los demás todavía tienen el
            # stock "inflado" de la carga histórica y darían alertas falsas.
            repuestos = [self._gestor.obtener_repuesto_id(i) for i in contados]
            resultado.alertas_nuevas = self._motor.revisar_todos(repuestos)
        return resultado
