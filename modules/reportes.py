"""Historial de movimientos y su descarga en Excel o PDF (Paso 10).

    FiltroMovimientos          -> qué se quiere ver (fechas, tipo, repuesto), con su contrato
    ConsultaMovimientos        -> busca en la base (solo lectura)
    ExportadorReporte (abstracta)
        ├── ExportadorExcel    -> .xlsx (openpyxl)
        └── ExportadorPDF      -> .pdf  (fpdf2)

Los exportadores son otra vez el patrón Strategy (Unidad 3): la pantalla pide
"un exportador" y le da lo mismo cuál es. Agregar CSV sería una clase más,
sin tocar la consulta ni la pantalla (abierto/cerrado).
"""
import io
import sqlite3
from abc import ABC, abstractmethod
from dataclasses import dataclass
from datetime import date
from typing import Dict, List, Optional

from modules.modelos import TIPOS_MOVIMIENTO

# Palabras del taller para cada tipo de movimiento (las mismas del formulario).
ETIQUETAS_TIPO_MOVIMIENTO = {
    "INGRESO": "Agregar (compra)",
    "EGRESO": "Quitar (uso)",
    "AJUSTE": "Conteo",
}

COLUMNAS_REPORTE = ("Fecha", "Movimiento", "Código", "Marca", "Categoría", "Cantidad", "Unidad", "Proveedor")


class FiltroInvalidoError(ValueError):
    """El filtro no tiene sentido (fechas inválidas, desde > hasta o tipo desconocido)."""


@dataclass(frozen=True)
class FiltroMovimientos:
    """Qué movimientos mostrar.

    Invariante: fechas ISO válidas, desde <= hasta y tipo en TIPOS_MOVIMIENTO (o None = todos).
    """
    desde: str
    hasta: str
    tipo: Optional[str] = None
    repuesto_id: Optional[int] = None

    def __post_init__(self):
        try:
            desde, hasta = date.fromisoformat(self.desde), date.fromisoformat(self.hasta)
        except (TypeError, ValueError):
            raise FiltroInvalidoError("Las fechas del filtro no son válidas.") from None
        if desde > hasta:
            raise FiltroInvalidoError("La fecha \"desde\" no puede ser posterior a \"hasta\".")
        if self.tipo is not None and self.tipo not in TIPOS_MOVIMIENTO:
            raise FiltroInvalidoError(f"Tipo de movimiento desconocido: {self.tipo}.")

    def descripcion(self) -> str:
        """Texto para el encabezado de los archivos: "Del 01/10/2026 al 20/10/2026 · Quitar (uso)"."""
        texto = f"Del {_fecha_ar(self.desde)} al {_fecha_ar(self.hasta)}"
        if self.tipo:
            texto += f" · {ETIQUETAS_TIPO_MOVIMIENTO[self.tipo]}"
        return texto


@dataclass(frozen=True)
class FilaReporte:
    """Un movimiento listo para mostrar (ya con nombres en vez de ids)."""
    id: int
    fecha: str
    tipo: str
    codigo: str
    marca: str
    categoria: str
    cantidad: float
    unidad_medida: str
    proveedor: str

    @property
    def tipo_texto(self) -> str:
        return ETIQUETAS_TIPO_MOVIMIENTO[self.tipo]

    @property
    def fecha_texto(self) -> str:
        return _fecha_ar(self.fecha)

    @property
    def cantidad_texto(self) -> str:
        """Ingresos con +, egresos con −, conteos con el signo de la corrección."""
        numero = f"{abs(self.cantidad):g}".replace(".", ",")
        if self.tipo == "INGRESO" or (self.tipo == "AJUSTE" and self.cantidad > 0):
            return f"+{numero}"
        return f"-{numero}"

    def como_tupla(self) -> tuple:
        """En el orden de COLUMNAS_REPORTE."""
        return (self.fecha_texto, self.tipo_texto, self.codigo, self.marca, self.categoria,
                self.cantidad_texto, self.unidad_medida, self.proveedor)


def _fecha_ar(iso: str) -> str:
    return date.fromisoformat(iso).strftime("%d/%m/%Y")


# ---------------------------------------------------------------- Consulta

class ConsultaMovimientos:
    """Busca movimientos con su repuesto, marca, categoría y proveedor. Solo lee."""

    def __init__(self, conexion: sqlite3.Connection):
        self.conexion = conexion

    def buscar(self, filtro: FiltroMovimientos) -> List[FilaReporte]:
        """Post: movimientos entre `desde` y `hasta` inclusive, del más nuevo al más viejo."""
        condiciones = ["m.fecha BETWEEN ? AND ?"]
        parametros: list = [filtro.desde, filtro.hasta]
        if filtro.tipo:
            condiciones.append("m.tipo = ?")
            parametros.append(filtro.tipo)
        if filtro.repuesto_id is not None:
            condiciones.append("m.repuesto_id = ?")
            parametros.append(filtro.repuesto_id)
        consulta = f"""
            SELECT m.id, m.fecha, m.tipo, m.cantidad,
                   r.codigo, r.unidad_medida,
                   ma.nombre AS marca, c.nombre AS categoria,
                   COALESCE(p.nombre, '') AS proveedor
            FROM movimientos_stock m
            JOIN repuestos r   ON r.id = m.repuesto_id
            JOIN marcas ma     ON ma.id = r.marca_id
            JOIN categorias c  ON c.id = r.categoria_id
            LEFT JOIN proveedores p ON p.id = m.proveedor_id
            WHERE {" AND ".join(condiciones)}
            ORDER BY m.fecha DESC, m.id DESC"""
        return [
            FilaReporte(id=f["id"], fecha=f["fecha"], tipo=f["tipo"], codigo=f["codigo"], marca=f["marca"],
                        categoria=f["categoria"], cantidad=f["cantidad"], unidad_medida=f["unidad_medida"],
                        proveedor=f["proveedor"])
            for f in self.conexion.execute(consulta, parametros).fetchall()
        ]


def contar_por_tipo(filas: List[FilaReporte]) -> Dict[str, int]:
    """Cuántos movimientos hay de cada tipo (no se suman cantidades: mezclan unidades y cm)."""
    conteo = {tipo: 0 for tipo in TIPOS_MOVIMIENTO}
    for fila in filas:
        conteo[fila.tipo] += 1
    return conteo


# ---------------------------------------------------------------- Exportadores (Strategy)

class ExportadorReporte(ABC):
    """Contrato: convierte las filas en los bytes de un archivo."""
    extension: str = ""
    tipo_mime: str = ""

    @abstractmethod
    def exportar(self, filas: List[FilaReporte], filtro: FiltroMovimientos) -> bytes:
        """Post: devuelve el archivo completo; una lista vacía da un archivo con solo el encabezado."""


class ExportadorExcel(ExportadorReporte):
    extension = "xlsx"
    tipo_mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    def exportar(self, filas: List[FilaReporte], filtro: FiltroMovimientos) -> bytes:
        from openpyxl import Workbook
        from openpyxl.styles import Font

        libro = Workbook()
        hoja = libro.active
        hoja.title = "Movimientos"
        hoja.append(("Movimientos de stock — Jannon & Durand",))
        hoja.append((filtro.descripcion(),))
        hoja.append(())
        hoja.append(COLUMNAS_REPORTE)
        hoja["A1"].font = Font(bold=True, size=13)
        for celda in hoja[4]:
            celda.font = Font(bold=True)
        for fila in filas:
            valores = list(fila.como_tupla())
            valores[5] = fila.cantidad if fila.tipo != "EGRESO" else -fila.cantidad  # número, para poder sumar en Excel
            hoja.append(valores)
        for letra, ancho in zip("ABCDEFGH", (12, 18, 26, 14, 20, 10, 9, 22)):
            hoja.column_dimensions[letra].width = ancho
        hoja.freeze_panes = "A5"
        salida = io.BytesIO()
        libro.save(salida)
        return salida.getvalue()


class ExportadorPDF(ExportadorReporte):
    extension = "pdf"
    tipo_mime = "application/pdf"
    _ANCHOS = (22, 30, 50, 24, 32, 18, 15, 66)  # mm, suman 257 (A4 apaisada: 297 - 2 márgenes de 15 mm = 267)

    def exportar(self, filas: List[FilaReporte], filtro: FiltroMovimientos) -> bytes:
        from fpdf import FPDF

        pdf = FPDF(orientation="L", format="A4")
        pdf.set_margins(15, 15, 15)
        pdf.set_auto_page_break(True, margin=15)
        pdf.add_page()
        pdf.set_font("Helvetica", "B", 14)
        pdf.cell(0, 8, _latin1("Movimientos de stock - Jannon & Durand"), new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", "", 10)
        pdf.cell(0, 6, _latin1(f"{filtro.descripcion()} · {len(filas)} movimientos"), new_x="LMARGIN", new_y="NEXT")
        pdf.ln(3)
        self._encabezado(pdf)
        pdf.set_font("Helvetica", "", 9)
        for fila in filas:
            if pdf.will_page_break(6):
                pdf.add_page()
                self._encabezado(pdf)
                pdf.set_font("Helvetica", "", 9)
            for valor, ancho in zip(fila.como_tupla(), self._ANCHOS):
                pdf.cell(ancho, 6, _latin1(_recortar(pdf, str(valor), ancho)), border="B")
            pdf.ln()
        if not filas:
            pdf.cell(0, 8, _latin1("No hay movimientos con este filtro."))
        return bytes(pdf.output())

    def _encabezado(self, pdf) -> None:
        pdf.set_font("Helvetica", "B", 9)
        pdf.set_fill_color(227, 238, 232)  # --pino-100 de la interfaz
        for titulo, ancho in zip(COLUMNAS_REPORTE, self._ANCHOS):
            pdf.cell(ancho, 7, _latin1(titulo), fill=True)
        pdf.ln()


def _latin1(texto: str) -> str:
    """Las fuentes básicas del PDF solo tienen caracteres latinos (acentos y ñ sí)."""
    return texto.replace("—", "-").replace("−", "-").encode("latin-1", "replace").decode("latin-1")


def _recortar(pdf, texto: str, ancho: float) -> str:
    """Corta el texto con "..." si no entra en la columna."""
    if pdf.get_string_width(_latin1(texto)) <= ancho - 2:
        return texto
    while texto and pdf.get_string_width(_latin1(texto + "...")) > ancho - 2:
        texto = texto[:-1]
    return texto + "..."


EXPORTADORES: Dict[str, ExportadorReporte] = {
    "excel": ExportadorExcel(),
    "pdf": ExportadorPDF(),
}
