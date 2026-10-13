"""Módulo de Ingesta y Migración de Datos Históricos (Proceso ETL).

Este módulo cumple la función exclusiva de extracción, transformación y carga para la migración inicial del historial del taller"""

import sqlite3
from datetime import datetime
 
from openpyxl import load_workbook
 
from modules.parsers import obtener_parser
 
HOJA_DATOS = "Productos"
ENCABEZADOS = ["Fecha", "N° Remito", "Proveedor", "Código", "Descripción", "Cantidad"]
 
 
def _fecha_iso(valor) -> str:
    """El Excel trae las fechas como texto dd/mm/aaaa; SQLite guarda mejor
    en formato ISO (comparable como texto y compatible con funciones de
    fecha de SQLite)."""
    if isinstance(valor, datetime):
        return valor.date().isoformat()
    return datetime.strptime(str(valor).strip(), "%d/%m/%Y").date().isoformat()
 
 
def _obtener_o_crear(conexion: sqlite3.Connection, tabla: str, columna: str, valor: str) -> int:
    fila = conexion.execute(f"SELECT id FROM {tabla} WHERE {columna} = ?", (valor,)).fetchone()
    if fila:
        return fila[0]
    cursor = conexion.execute(f"INSERT INTO {tabla} ({columna}) VALUES (?)", (valor,))
    return cursor.lastrowid
 
 
def _obtener_o_crear_repuesto(conexion: sqlite3.Connection, linea) -> int:
    """Resuelve la identidad del repuesto por (marca, código): el código es
    el del fabricante y no depende de quién lo vende, así que el mismo
    repuesto comprado a dos proveedores queda como UN solo repuesto.
    Si ya existe esa combinación, lo reutiliza; si no, lo crea con su
    categoría (ya unificada por el parser) y sus especificaciones.
    """
    marca_id = _obtener_o_crear(conexion, "marcas", "nombre", linea.marca)
    fila = conexion.execute(
        "SELECT id FROM repuestos WHERE marca_id = ? AND codigo = ?",
        (marca_id, linea.codigo),
    ).fetchone()
    if fila:
        return fila[0]

    categoria_id = _obtener_o_crear(conexion, "categorias", "nombre", linea.categoria)
    cursor = conexion.execute(
        """INSERT INTO repuestos (codigo, marca_id, categoria_id, especificaciones)
           VALUES (?, ?, ?, ?)""",
        (linea.codigo, marca_id, categoria_id, linea.especificaciones),
    )
    return cursor.lastrowid
 
 
def cargar_historico(conexion: sqlite3.Connection, ruta_excel: str) -> None:
    """Punto de entrada del importador. Asume que el esquema ya existe
    (ver `modules.base_datos.crear_esquema`).
 
    No hace commit ni rollback propios: la transacción la controla quien
    llama (`modules.base_datos.inicializar_si_hace_falta`), para poder
    deshacer todo de una si algo falla a mitad de la carga.
    """
    libro = load_workbook(ruta_excel, read_only=True, data_only=True)
    hoja = libro[HOJA_DATOS]
 
    total_producto = 0
    total_cargo_excluido = 0
 
    for valores in hoja.iter_rows(min_row=2, values_only=True):
        if valores[0] is None:  
            continue
        fila = dict(zip(ENCABEZADOS, valores))
 
        parser = obtener_parser(fila["Proveedor"])
        linea = parser.parsear(fila)
 
        if linea.tipo_linea != "PRODUCTO":
            total_cargo_excluido += 1
            continue
 
        proveedor_id = _obtener_o_crear(conexion, "proveedores", "nombre", linea.proveedor)
        repuesto_id = _obtener_o_crear_repuesto(conexion, linea)
        fecha = _fecha_iso(fila["Fecha"])
 
        conexion.execute(
            """INSERT INTO movimientos_stock
               (repuesto_id, proveedor_id, fecha, tipo, cantidad, codigo_proveedor_raw, origen)
               VALUES (?, ?, ?, 'INGRESO', ?, ?, 'carga_inicial')""",
            (repuesto_id, proveedor_id, fecha, linea.cantidad, linea.codigo),
        )
        total_producto += 1
 
    print(
        f"Carga inicial completa: {total_producto} movimientos de stock cargados, "
        f"{total_cargo_excluido} líneas de cargos administrativos excluidas (no son stock)."
    )