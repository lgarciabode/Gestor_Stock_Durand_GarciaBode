"""Pruebas del conteo inicial con planilla (Paso 7 del modulo de alertas).

- PlanillaConteo: genera y lee el Excel (archivo temporal).
- ServicioConteoInicial: todo o nada, validaciones por fila y flujo completo
  (ajustes -> umbrales -> alertas) sobre una base SQLite temporal.
- Un Mock de la fabrica de parametros verifica COMO se usa ese colaborador
  (Unidad 4, dobles de prueba).
- Integracion con el Excel historico real (192 repuestos).

unittest puro, patron AAA (Arrange, Act, Assert).
"""
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import load_workbook

from modules.alertas import MotorAlertas, reglas_por_defecto
from modules.base_datos import crear_esquema, obtener_conexion
from modules.carga_historico_datos import cargar_historico
from modules.conteo import (
    COLUMNAS,
    LecturaPlanilla,
    FilaConteo,
    PlanillaConteo,
    PlanillaConteoInvalidaError,
    ServicioConteoInicial,
)
from modules.crud import GestorInventarioBD, RepositorioAlertas
from modules.modelos import Categoria, Marca, MovimientoStock, Repuesto

RUTA_EXCEL = Path(__file__).resolve().parent.parent / "data" / "base_stock_taller.xlsx"
HOY = "2026-10-20"


def completar(ruta, cantidades_por_codigo):
    """Simula a la persona que completa la columna "Cantidad contada"."""
    libro = load_workbook(ruta)
    hoja = libro.active
    for fila in hoja.iter_rows(min_row=2):
        if fila[1].value in cantidades_por_codigo:
            fila[7].value = cantidades_por_codigo[fila[1].value]
    libro.save(ruta)


class BaseTemporal(unittest.TestCase):
    """Dos rodamientos ZKL y un tubo SAV, con compras cargadas."""

    def setUp(self):
        self.directorio = tempfile.TemporaryDirectory()
        self.ruta = Path(self.directorio.name) / "conteo.xlsx"
        self.conexion = obtener_conexion(str(Path(self.directorio.name) / "test.sqlite3"))
        crear_esquema(self.conexion)
        self.gestor = GestorInventarioBD(self.conexion)
        zkl = self.gestor.crear_marca(Marca(id=None, nombre="ZKL"))
        sav = self.gestor.crear_marca(Marca(id=None, nombre="SAV"))
        rod = self.gestor.crear_categoria(Categoria(id=None, nombre="RODAMIENTO"))
        tubo = self.gestor.crear_categoria(Categoria(id=None, nombre="TUBO"))
        self.ids = {
            "6204 2RS": self.gestor.crear_repuestos(Repuesto(id=None, codigo="6204 2RS", marca_id=zkl, categoria_id=rod)),
            "6205 2RS": self.gestor.crear_repuestos(Repuesto(id=None, codigo="6205 2RS", marca_id=zkl, categoria_id=rod)),
            "T-40": self.gestor.crear_repuestos(Repuesto(id=None, codigo="T-40", marca_id=sav, categoria_id=tubo)),
        }
        for codigo, fechas in (("6204 2RS", [("2026-02-13", 10), ("2026-04-23", 15), ("2026-06-30", 20), ("2026-08-05", 10)]),
                               ("6205 2RS", [("2026-03-01", 6)]),
                               ("T-40", [("2026-01-05", 4)])):
            for fecha, cantidad in fechas:
                self.gestor.registrar_movimiento(MovimientoStock(
                    id=None, repuesto_id=self.ids[codigo], fecha=fecha, tipo="INGRESO", cantidad=cantidad))
        self.planilla = PlanillaConteo()
        self.motor = MotorAlertas(reglas_por_defecto(), RepositorioAlertas(self.conexion), hoy=lambda: HOY)
        self.servicio = ServicioConteoInicial(self.gestor, self.motor)

    def tearDown(self):
        self.conexion.close()
        self.directorio.cleanup()

    def generar(self):
        marcas = {m.id: m.nombre for m in self.gestor.obtener_marcas()}
        categorias = {c.id: c.nombre for c in self.gestor.obtener_categorias()}
        return self.planilla.generar(self.ruta, self.gestor.obtener_repuestos(), marcas, categorias)

    def stock(self, codigo):
        return self.gestor.obtener_repuesto_id(self.ids[codigo]).stock_actual


# ================================================================ Planilla

class PruebaPlanilla(BaseTemporal):

    def test_genera_un_renglon_por_repuesto_ordenado_para_recorrer_el_taller(self):
        # Act
        cantidad = self.generar()

        # Assert: orden categoria -> marca -> codigo; la columna a completar vacia
        hoja = load_workbook(self.ruta).active
        filas = list(hoja.iter_rows(values_only=True))
        self.assertEqual(cantidad, 3)
        self.assertEqual(filas[0], COLUMNAS)
        self.assertEqual([f[1] for f in filas[1:]], ["6204 2RS", "6205 2RS", "T-40"])
        self.assertEqual(filas[1][6], 55)
        self.assertTrue(all(f[7] is None for f in filas[1:]))

    def test_la_columna_a_completar_solo_acepta_numeros(self):
        self.generar()
        validaciones = load_workbook(self.ruta).active.data_validations.dataValidation
        self.assertEqual(len(validaciones), 1)
        self.assertEqual(str(validaciones[0].sqref), "H2:H4")

    def test_leer_lo_completado(self):
        # Arrange
        self.generar()
        completar(self.ruta, {"6204 2RS": 3, "T-40": "2,5"})

        # Act
        lectura = self.planilla.leer(self.ruta)

        # Assert: coma decimal como texto se acepta; vacio = sin contar
        self.assertEqual(lectura.errores, [])
        self.assertEqual([(f.codigo, f.cantidad_contada) for f in lectura.filas],
                         [("6204 2RS", 3.0), ("6205 2RS", None), ("T-40", 2.5)])
        self.assertEqual(lectura.filas[0].fila, 2)

    def test_texto_en_la_cantidad_es_un_error_con_su_fila(self):
        self.generar()
        completar(self.ruta, {"6205 2RS": "tres"})

        lectura = self.planilla.leer(self.ruta)

        self.assertEqual(len(lectura.errores), 1)
        self.assertEqual((lectura.errores[0].fila, lectura.errores[0].codigo), (3, "6205 2RS"))

    def test_un_archivo_que_no_es_la_planilla_se_rechaza(self):
        # Arrange: le cambiaron un titulo
        self.generar()
        libro = load_workbook(self.ruta)
        libro.active["H1"] = "Cuantos hay"
        libro.save(self.ruta)

        with self.assertRaises(PlanillaConteoInvalidaError):
            self.planilla.leer(self.ruta)


# ================================================================ Servicio

class PruebaServicioConteo(BaseTemporal):

    def test_todo_o_nada_si_una_fila_esta_mal_no_se_guarda_ninguna(self):
        # Arrange: una fila bien y otra con decimales en un repuesto por unidad
        self.generar()
        completar(self.ruta, {"6204 2RS": 3, "6205 2RS": 2.5})

        # Act
        resultado = self.servicio.aplicar(self.planilla.leer(self.ruta), HOY)

        # Assert
        self.assertFalse(resultado.aplicado)
        self.assertEqual([(e.fila, e.codigo) for e in resultado.errores], [(3, "6205 2RS")])
        self.assertEqual(self.stock("6204 2RS"), 55)  # no se toco

    def test_codigo_que_no_corresponde_al_id(self):
        lectura = LecturaPlanilla(filas=[FilaConteo(2, self.ids["6204 2RS"], "T-40", 1)])

        errores = self.servicio.validar(lectura.filas, HOY)

        self.assertIn('corresponde al código "6204 2RS"', errores[0].mensaje)

    def test_id_repetido_e_id_inexistente(self):
        filas = [FilaConteo(2, self.ids["T-40"], "T-40", 1),
                 FilaConteo(3, self.ids["T-40"], "T-40", 2),
                 FilaConteo(4, 999, "X", 1)]

        errores = self.servicio.validar(filas, HOY)

        self.assertEqual([e.fila for e in errores], [3, 4])

    def test_conteo_parcial_corrige_lo_contado_y_saltea_lo_vacio(self):
        # Arrange
        self.generar()
        completar(self.ruta, {"6204 2RS": 3, "T-40": 4})

        # Act
        resultado = self.servicio.aplicar(self.planilla.leer(self.ruta), HOY)

        # Assert
        self.assertTrue(resultado.aplicado)
        self.assertEqual((resultado.ajustados, resultado.sin_diferencia, resultado.sin_contar), (1, 1, 1))
        self.assertEqual(self.stock("6204 2RS"), 3)
        self.assertEqual(self.stock("6205 2RS"), 6)  # no se conto: queda igual

    def test_despues_del_conteo_recalcula_umbrales_y_revisa_alertas(self):
        self.generar()
        completar(self.ruta, {"6204 2RS": 3, "6205 2RS": 6, "T-40": 4})

        resultado = self.servicio.aplicar(self.planilla.leer(self.ruta), HOY)

        self.assertEqual(resultado.umbrales_recalculados, 3)
        rep = self.gestor.obtener_repuesto_id(self.ids["6204 2RS"])
        self.assertTrue(rep.tiene_umbrales())
        tipos = {(a.repuesto_id, a.tipo) for a in resultado.alertas_nuevas}
        self.assertIn((self.ids["6204 2RS"], "REORDEN"), tipos)

    def test_en_un_conteo_parcial_no_hay_alertas_de_lo_que_no_se_conto(self):
        # Arrange: al T-40 se le ponen umbrales a mano (maximo 2) con 4 en stock:
        # si se revisara daria "Sobra stock", pero no se conto
        self.gestor.actualizar_parametros_stock(self.ids["T-40"], stock_minimo=0, stock_maximo=2, punto_reorden=1)
        self.assertTrue(self.gestor.obtener_repuesto_id(self.ids["T-40"]).tiene_sobrestock())
        self.generar()
        completar(self.ruta, {"6204 2RS": 3})

        # Act
        resultado = self.servicio.aplicar(self.planilla.leer(self.ruta), HOY)

        # Assert: solo la alerta del repuesto contado
        self.assertEqual({a.repuesto_id for a in resultado.alertas_nuevas}, {self.ids["6204 2RS"]})

    def test_los_umbrales_se_recalculan_despues_de_guardar_los_ajustes(self):
        # Mock: se verifica que la fabrica se usa una vez, con el gestor,
        # cuando el ajuste ya esta guardado (el periodo llega hasta HOY)
        fabrica = Mock()
        fabrica.return_value.recalcular_todos.side_effect = lambda: self.assertEqual(self.stock("6204 2RS"), 3) or 0
        servicio = ServicioConteoInicial(self.gestor, self.motor, fabrica_parametros=fabrica)

        servicio.aplicar(LecturaPlanilla(filas=[FilaConteo(2, self.ids["6204 2RS"], "6204 2RS", 3)]), HOY)

        fabrica.assert_called_once_with(self.gestor)
        fabrica.return_value.recalcular_todos.assert_called_once_with()

    def test_con_errores_no_recalcula_nada(self):
        fabrica = Mock()
        servicio = ServicioConteoInicial(self.gestor, self.motor, fabrica_parametros=fabrica)

        servicio.aplicar(LecturaPlanilla(filas=[FilaConteo(2, 999, "X", 1)]), HOY)

        fabrica.assert_not_called()

    def test_planilla_vacia_no_recalcula_nada(self):
        fabrica = Mock()
        servicio = ServicioConteoInicial(self.gestor, self.motor, fabrica_parametros=fabrica)
        self.generar()

        resultado = servicio.aplicar(self.planilla.leer(self.ruta), HOY)

        self.assertEqual(resultado.sin_contar, 3)
        fabrica.assert_not_called()


@unittest.skipUnless(RUTA_EXCEL.exists(), "No está el Excel histórico en data/")
class PruebaConteoInicialSobreDatosReales(unittest.TestCase):

    def setUp(self):
        self.directorio = tempfile.TemporaryDirectory()
        self.ruta = Path(self.directorio.name) / "conteo.xlsx"
        self.conexion = obtener_conexion(str(Path(self.directorio.name) / "test.sqlite3"))
        crear_esquema(self.conexion)
        cargar_historico(self.conexion, str(RUTA_EXCEL))
        self.conexion.commit()
        self.gestor = GestorInventarioBD(self.conexion)

    def tearDown(self):
        self.conexion.close()
        self.directorio.cleanup()

    def test_planilla_de_los_192_y_conteo_del_6204(self):
        # Arrange
        marcas = {m.id: m.nombre for m in self.gestor.obtener_marcas()}
        categorias = {c.id: c.nombre for c in self.gestor.obtener_categorias()}
        cantidad = PlanillaConteo().generar(self.ruta, self.gestor.obtener_repuestos(), marcas, categorias)
        completar(self.ruta, {"6204 2RS": 3})
        motor = MotorAlertas(reglas_por_defecto(), RepositorioAlertas(self.conexion), hoy=lambda: HOY)

        # Act
        resultado = ServicioConteoInicial(self.gestor, motor).aplicar(PlanillaConteo().leer(self.ruta), HOY)

        # Assert
        self.assertEqual(cantidad, 192)
        self.assertEqual((resultado.ajustados, resultado.sin_contar), (1, 191))
        self.assertEqual(resultado.umbrales_recalculados, 192)
        rep = next(r for r in self.gestor.obtener_repuestos() if r.codigo == "6204 2RS")
        self.assertEqual(rep.stock_actual, 3)


if __name__ == "__main__":
    unittest.main()
