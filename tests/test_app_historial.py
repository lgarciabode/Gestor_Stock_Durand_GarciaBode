"""Pruebas de la pantalla "Movimientos" y sus descargas (Paso 10).

Cliente de prueba de Flask sobre una base SQLite temporal. Si Flask o
auth-module no están instalados, se saltean (y la del PDF, si falta fpdf2).

unittest puro, patron AAA (Arrange, Act, Assert).
"""
import importlib.util
import io
import os
import sys
import tempfile
import unittest
from datetime import date, timedelta
from pathlib import Path
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

HAY_FLASK = all(importlib.util.find_spec(m) for m in ("flask", "auth_module"))
HAY_FPDF = importlib.util.find_spec("fpdf") is not None

if HAY_FLASK:
    os.environ.setdefault("DATABASE_URL", "data/stock_taller.sqlite3")  # config.py lo exige al importarse
    from apps import app_1

from openpyxl import load_workbook

from modules.base_datos import crear_esquema, obtener_conexion
from modules.crud import GestorInventarioBD
from modules.modelos import Categoria, Marca, MovimientoStock, Repuesto

HOY = date.today()
HACE_10 = (HOY - timedelta(days=10)).isoformat()
HACE_60 = (HOY - timedelta(days=60)).isoformat()


@unittest.skipUnless(HAY_FLASK, "Flask / auth-module no instalados")
class PruebaPantallaHistorial(unittest.TestCase):

    def setUp(self):
        self.directorio = tempfile.TemporaryDirectory()
        ruta_bd = str(Path(self.directorio.name) / "stock.sqlite3")
        conexion = obtener_conexion(ruta_bd)
        crear_esquema(conexion)
        conexion.execute("UPDATE estado_sistema SET carga_inicial_completa = 1 WHERE id = 1")
        gestor = GestorInventarioBD(conexion)
        zkl = gestor.crear_marca(Marca(id=None, nombre="ZKL"))
        rod = gestor.crear_categoria(Categoria(id=None, nombre="RODAMIENTO"))
        a = gestor.crear_repuestos(Repuesto(id=None, codigo="6204 2RS", marca_id=zkl, categoria_id=rod))
        b = gestor.crear_repuestos(Repuesto(id=None, codigo="6205 2RS", marca_id=zkl, categoria_id=rod))
        for rid, fecha, tipo, cantidad in ((a, HACE_60, "INGRESO", 20), (a, HACE_10, "EGRESO", 3), (b, HACE_10, "INGRESO", 4)):
            gestor.registrar_movimiento(MovimientoStock(id=None, repuesto_id=rid, fecha=fecha, tipo=tipo, cantidad=cantidad))
        conexion.close()

        self.parches = [
            patch.object(app_1, "DATABASE_URL", ruta_bd),
            patch.object(app_1, "SECRET_KEY", "clave-de-prueba"),
            patch.object(app_1, "RUTA_BD_USUARIOS", str(Path(self.directorio.name) / "usuarios.sqlite3")),
        ]
        for p in self.parches:
            p.start()
        app = app_1.create_app()
        app.config["TESTING"] = True
        self.cliente = app.test_client()
        with self.cliente.session_transaction() as sesion:
            sesion["email"] = "cami@taller.com"

    def tearDown(self):
        for p in self.parches:
            p.stop()
        self.directorio.cleanup()

    def html(self, url):
        return self.cliente.get(url).get_data(as_text=True)

    # ------------------------------------------------------------ pantalla

    def test_por_defecto_muestra_los_ultimos_30_dias(self):
        html = self.html("/movimientos")
        self.assertIn("2 movimientos", html)          # el de hace 60 dias no aparece
        self.assertIn("Quitar (uso): 1", html)

    def test_filtrar_por_fechas_tipo_y_repuesto(self):
        self.assertIn("3 movimientos", self.html(f"/movimientos?desde={HACE_60}"))
        self.assertIn("1 movimiento<", self.html(f"/movimientos?desde={HACE_60}&tipo=EGRESO"))
        self.assertIn("2 movimientos", self.html(f"/movimientos?desde={HACE_60}&repuesto=6204+2RS"))

    def test_filtro_invalido_muestra_el_motivo(self):
        respuesta = self.cliente.get(f"/movimientos?desde={HOY.isoformat()}&hasta={HACE_10}")
        self.assertEqual(respuesta.status_code, 400)
        self.assertIn("no puede ser posterior", respuesta.get_data(as_text=True))

    def test_los_links_de_descarga_repiten_el_filtro(self):
        html = self.html(f"/movimientos?desde={HACE_60}&tipo=EGRESO")
        self.assertIn(f"/movimientos/descargar/excel?desde={HACE_60}", html)
        self.assertIn("tipo=EGRESO", html)

    # ------------------------------------------------------------ descargas

    def test_descargar_excel_con_lo_filtrado(self):
        # Act
        respuesta = self.cliente.get(f"/movimientos/descargar/excel?desde={HACE_60}&tipo=INGRESO")

        # Assert
        self.assertEqual(respuesta.status_code, 200)
        self.assertIn("attachment", respuesta.headers["Content-Disposition"])
        self.assertIn(".xlsx", respuesta.headers["Content-Disposition"])
        hoja = load_workbook(io.BytesIO(respuesta.data)).active
        self.assertEqual(hoja.max_row, 4 + 2)  # encabezado + 2 compras

    @unittest.skipUnless(HAY_FPDF, "fpdf2 no instalado")
    def test_descargar_pdf(self):
        respuesta = self.cliente.get("/movimientos/descargar/pdf")
        self.assertEqual(respuesta.mimetype, "application/pdf")
        self.assertTrue(respuesta.data.startswith(b"%PDF"))

    def test_formato_desconocido(self):
        self.assertEqual(self.cliente.get("/movimientos/descargar/word").status_code, 404)

    def test_descarga_con_filtro_invalido_vuelve_a_la_pantalla(self):
        respuesta = self.cliente.get("/movimientos/descargar/excel?desde=ayer", follow_redirects=True)
        self.assertEqual(respuesta.request.path, "/movimientos")
        self.assertIn("no son válidas", respuesta.get_data(as_text=True))


if __name__ == "__main__":
    unittest.main()
